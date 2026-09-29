from dataclasses import dataclass
from datetime import timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.core.errors import bad_request, not_found
from app.models import ExportLigne, ExportLot, Pige, Utilisateur
from app.schemas import ExportLigneRead, ExportLotRead
from app.services.besoins import get_besoins_rows, ordre_coffret, valeur_code


def _lot_query():
    return select(ExportLot).options(
        selectinload(ExportLot.utilisateur),
        selectinload(ExportLot.lignes).selectinload(ExportLigne.pige).selectinload(Pige.coffret),
    )


def _lot_to_read(lot: ExportLot) -> ExportLotRead:
    lignes = sorted(
        lot.lignes,
        key=lambda ligne: (ligne.pige.coffret.nom, ligne.pige.position_ligne, ligne.pige.position_colonne),
    )
    return ExportLotRead(
        id=lot.id,
        date=lot.date if lot.date.tzinfo else lot.date.replace(tzinfo=timezone.utc),
        utilisateur_nom=lot.utilisateur.nom if lot.utilisateur else "Non renseigne",
        lignes=[
            ExportLigneRead(
                pige_id=ligne.pige_id,
                coffret_nom=ligne.pige.coffret.nom,
                code=ligne.pige.code,
                quantite=ligne.quantite,
            )
            for ligne in lignes
        ],
    )


def create_export_lot(db: Session, utilisateur_id: int) -> ExportLotRead:
    utilisateur = db.get(Utilisateur, utilisateur_id)
    if utilisateur is None:
        raise not_found("Utilisateur introuvable")

    besoins = get_besoins_rows(db)
    if not besoins:
        raise bad_request("Aucune nouvelle piece a exporter")

    lot = ExportLot(
        utilisateur_id=utilisateur.id,
        lignes=[ExportLigne(pige_id=row.pige_id, quantite=row.quantite_a_commander) for row in besoins],
    )
    db.add(lot)
    db.commit()
    return get_export_lot(db, lot.id)


def list_export_lots(db: Session) -> list[ExportLotRead]:
    lots = db.scalars(_lot_query().order_by(ExportLot.id.desc())).all()
    return [_lot_to_read(lot) for lot in lots]


def get_export_lot(db: Session, lot_id: int) -> ExportLotRead:
    lot = db.scalar(_lot_query().where(ExportLot.id == lot_id))
    if lot is None:
        raise not_found("Export introuvable")
    return _lot_to_read(lot)


@dataclass
class LigneCommande:
    code: str
    quantite: int
    coffrets: list[str]


def lignes_par_code(lot: ExportLotRead) -> list[LigneCommande]:
    """Une ligne par code de pige : les quantites de tous les coffrets sont additionnees."""
    lignes: dict[str, LigneCommande] = {}
    for ligne in lot.lignes:
        commande = lignes.setdefault(ligne.code, LigneCommande(ligne.code, 0, []))
        commande.quantite += ligne.quantite
        nom = ligne.coffret_nom.capitalize()
        if nom not in commande.coffrets:
            commande.coffrets.append(nom)
    for commande in lignes.values():
        commande.coffrets.sort(key=ordre_coffret)
    return sorted(lignes.values(), key=lambda commande: valeur_code(commande.code))


def _entete(lot: ExportLotRead) -> tuple[str, str]:
    return (
        f"PigeControl - Commande n°{lot.id}",
        lot.date.astimezone().strftime("Export généré le %d/%m/%Y à %H:%M") + f" par {lot.utilisateur_nom}",
    )


def build_lot_excel(lot: ExportLotRead) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Commande"

    titre, sous_titre = _entete(lot)
    sheet.append([titre])
    sheet.append([sous_titre])
    sheet.append([])
    sheet["A1"].font = Font(bold=True, size=14, color="004E68")

    sheet.append(["Code pige", "Quantité à commander", "Coffrets"])
    header_fill = PatternFill("solid", fgColor="D9F3FB")
    for cell in sheet[4]:
        cell.font = Font(bold=True, color="004E68")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    lignes = lignes_par_code(lot)
    for ligne in lignes:
        sheet.append([ligne.code.replace(",", "."), ligne.quantite, ", ".join(ligne.coffrets)])
        sheet.cell(row=sheet.max_row, column=2).alignment = Alignment(horizontal="center")

    sheet.append(["Total", sum(ligne.quantite for ligne in lignes), ""])
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
    sheet.cell(row=sheet.max_row, column=2).alignment = Alignment(horizontal="center")

    sheet.column_dimensions["A"].width = 14
    sheet.column_dimensions["B"].width = 22
    sheet.column_dimensions["C"].width = 60
    sheet.freeze_panes = "A5"

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def build_lot_pdf(lot: ExportLotRead) -> BytesIO:
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36,
    )
    styles = getSampleStyleSheet()
    titre, sous_titre = _entete(lot)
    story = [
        Paragraph(titre, styles["Title"]),
        Paragraph(sous_titre, styles["Normal"]),
        Spacer(1, 18),
    ]

    lignes = lignes_par_code(lot)
    table_data = [["Code pige", "Quantité à commander", "Coffrets"]]
    for ligne in lignes:
        # Paragraph permet aux longues listes de coffrets de passer a la ligne.
        table_data.append([ligne.code.replace(",", "."), str(ligne.quantite), Paragraph(", ".join(ligne.coffrets), styles["Normal"])])
    if not lignes:
        table_data.append(["-", "0", "Aucun besoin"])
    table_data.append(["Total", str(sum(ligne.quantite for ligne in lignes)), ""])

    table = Table(table_data, colWidths=[80, 120, 323], repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#006D8F")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E6EEF2")),
                ("ALIGN", (0, 0), (1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD7DC")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F4F7F8")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 7),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ]
        )
    )
    story.append(table)
    document.build(story)
    output.seek(0)
    return output
