from __future__ import annotations

import argparse
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select

from app.database import Base, SessionLocal, engine
from app.models import Coffret, Pige

COFFRET_MARKER = "COFFRET PIGES"
DEFAULT_EXCEL_NAMES = ("Reappro PIGES.xlsm", "Réappro PIGES.xlsm")


@dataclass(frozen=True)
class CoffretBlock:
    sheet_name: str
    nom: str
    header_row: int
    start_col: int
    end_col: int


@dataclass(frozen=True)
class PigeRow:
    coffret_nom: str
    code: str
    quantite_initiale: int
    position_ligne: int
    position_colonne: int


def normalize_text(value: Any) -> str:
    return str(value).replace("\n", " ").strip()


def is_coffret_title(value: Any) -> bool:
    if value is None:
        return False
    return COFFRET_MARKER in normalize_text(value).upper()


def normalize_code(value: Any) -> str | None:
    if value is None:
        return None

    if isinstance(value, str):
        value = value.strip()
        if not value:
            return None
        value = value.replace(",", ".")

    try:
        decimal_value = Decimal(str(value))
    except (InvalidOperation, ValueError):
        return None

    if decimal_value < 0:
        return None

    decimal_value = decimal_value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    return f"{decimal_value:.2f}".replace(".", ",")


def normalize_quantity(value: Any) -> int:
    if value is None:
        return 1

    if isinstance(value, str):
        value = value.strip()
        if not value:
            return 1
        value = value.replace(",", ".")

    try:
        quantity = int(Decimal(str(value)))
    except (InvalidOperation, ValueError):
        return 1

    return quantity if quantity > 0 else 1


def get_merged_bounds_for_cell(sheet: Worksheet, row: int, col: int) -> tuple[int, int]:
    for merged_range in sheet.merged_cells.ranges:
        if (row, col) in merged_range.cells:
            return merged_range.min_col, merged_range.max_col

    return col, infer_unmerged_title_end_col(sheet, row, col)


def infer_unmerged_title_end_col(sheet: Worksheet, row: int, start_col: int) -> int:
    end_col = start_col
    empty_count = 0

    for col in range(start_col + 1, sheet.max_column + 1):
        value = sheet.cell(row=row + 1, column=col).value
        if normalize_code(value) is None:
            empty_count += 1
            if empty_count >= 2:
                break
            continue

        empty_count = 0
        end_col = col

    return end_col


def find_coffret_blocks(sheet: Worksheet) -> list[CoffretBlock]:
    blocks: list[CoffretBlock] = []

    for row in sheet.iter_rows():
        for cell in row:
            if not is_coffret_title(cell.value):
                continue

            start_col, end_col = get_merged_bounds_for_cell(sheet, cell.row, cell.column)
            blocks.append(
                CoffretBlock(
                    sheet_name=sheet.title,
                    nom=normalize_text(cell.value),
                    header_row=cell.row,
                    start_col=start_col,
                    end_col=end_col,
                )
            )

    return blocks


def row_code_count(dataframe: pd.DataFrame, row: int, start_col: int, end_col: int) -> int:
    count = 0
    for col in range(start_col, end_col + 1):
        if normalize_code(dataframe.iat[row - 1, col - 1]) is not None:
            count += 1
    return count


def extract_piges_from_block(dataframe: pd.DataFrame, block: CoffretBlock) -> list[PigeRow]:
    piges: list[PigeRow] = []
    position_ligne = 0
    empty_rows_after_grid = 0
    row = block.header_row + 1

    while row <= len(dataframe.index):
        code_count = row_code_count(dataframe, row, block.start_col, block.end_col)

        if code_count == 0:
            if piges:
                empty_rows_after_grid += 1
                if empty_rows_after_grid >= 2:
                    break
            row += 1
            continue

        empty_rows_after_grid = 0

        if code_count >= 2:
            position_ligne += 1
            for col in range(block.start_col, block.end_col + 1):
                code = normalize_code(dataframe.iat[row - 1, col - 1])
                if code is None:
                    continue

                quantity_value = None
                if row < len(dataframe.index):
                    quantity_value = dataframe.iat[row, col - 1]

                piges.append(
                    PigeRow(
                        coffret_nom=block.nom,
                        code=code,
                        quantite_initiale=normalize_quantity(quantity_value),
                        position_ligne=position_ligne,
                        position_colonne=col - block.start_col + 1,
                    )
                )

            row += 2
            continue

        row += 1

    return piges


def resolve_excel_path(raw_path: str | None) -> Path:
    candidates: list[Path] = []

    if raw_path:
        candidates.append(Path(raw_path))
    else:
        candidates.extend(Path(name) for name in DEFAULT_EXCEL_NAMES)
        candidates.extend(Path("..") / name for name in DEFAULT_EXCEL_NAMES)

    for candidate in candidates:
        resolved = candidate.expanduser().resolve()
        if resolved.exists():
            return resolved

    searched = ", ".join(str(candidate) for candidate in candidates)
    raise FileNotFoundError(f"Fichier Excel introuvable. Chemins testes : {searched}")


def import_piges(excel_path: Path) -> tuple[int, int]:
    workbook = load_workbook(excel_path, data_only=True, keep_vba=True)
    Base.metadata.create_all(bind=engine)

    imported_coffrets = 0
    imported_piges = 0

    with SessionLocal() as db:
        for sheet in workbook.worksheets:
            dataframe = pd.read_excel(
                excel_path,
                sheet_name=sheet.title,
                header=None,
                engine="openpyxl",
            )
            blocks = find_coffret_blocks(sheet)

            for block in blocks:
                coffret = db.scalar(select(Coffret).where(Coffret.nom == block.nom))
                if coffret is None:
                    coffret = Coffret(nom=block.nom)
                    db.add(coffret)
                    db.flush()
                    imported_coffrets += 1

                for pige_data in extract_piges_from_block(dataframe, block):
                    pige = db.scalar(
                        select(Pige).where(
                            Pige.coffret_id == coffret.id,
                            Pige.code == pige_data.code,
                        )
                    )

                    if pige is None:
                        pige = Pige(coffret_id=coffret.id, code=pige_data.code)
                        db.add(pige)

                    pige.quantite_initiale = pige_data.quantite_initiale
                    pige.position_ligne = pige_data.position_ligne
                    pige.position_colonne = pige_data.position_colonne
                    imported_piges += 1

        db.commit()

    return imported_coffrets, imported_piges


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Importer les coffrets et piges Excel dans SQLite."
    )
    parser.add_argument(
        "excel_path",
        nargs="?",
        help="Chemin vers le fichier Excel .xlsm. Exemple: '../Réappro PIGES.xlsm'",
    )
    args = parser.parse_args()

    excel_path = resolve_excel_path(args.excel_path)
    coffrets_count, piges_count = import_piges(excel_path)

    print(f"Import termine depuis : {excel_path}")
    print(f"Coffrets ajoutes : {coffrets_count}")
    print(f"Piges importees ou mises a jour : {piges_count}")


if __name__ == "__main__":
    main()
